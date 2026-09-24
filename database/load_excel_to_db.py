"""Importe l'onglet Données NBA dans la base SQLite définie par schema.sql.

Exemple : python database/load_excel_to_db.py "inputs/regular NBA.xlsx" --season 2024-25
Dépendances : openpyxl, pydantic (v2).
"""
from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, HttpUrl, ValidationError, create_model


# Colonnes source dans l'ordre de l'onglet « Données NBA » (A:AS).
STAT_COLUMNS = {
    4: 'games_played', 5: 'wins', 6: 'losses', 7: 'minutes_per_game',
    8: 'points_total', 9: 'field_goals_made', 10: 'field_goals_attempted',
    11: 'field_goal_pct', 12: 'three_points_made', 13: 'three_points_attempted',
    14: 'three_point_pct', 15: 'free_throws_made', 16: 'free_throws_attempted',
    17: 'free_throw_pct', 18: 'offensive_rebounds', 19: 'defensive_rebounds',
    20: 'total_rebounds', 21: 'assists', 22: 'turnovers', 23: 'steals',
    24: 'blocks', 25: 'personal_fouls', 26: 'fantasy_points',
    27: 'double_doubles', 28: 'triple_doubles', 29: 'plus_minus',
    30: 'offensive_rating', 31: 'defensive_rating', 32: 'net_rating',
    33: 'assist_pct', 34: 'assist_turnover_ratio', 35: 'assist_ratio',
    36: 'offensive_rebound_pct', 37: 'defensive_rebound_pct',
    38: 'total_rebound_pct', 39: 'turnover_ratio',
    40: 'effective_field_goal_pct', 41: 'true_shooting_pct',
    42: 'usage_pct', 43: 'pace', 44: 'player_impact_estimate', 45: 'possessions',
}
INTEGER_STATS = {'games_played', 'wins', 'losses', 'points_total',
                 'field_goals_made', 'field_goals_attempted', 'three_points_made',
                 'three_points_attempted', 'free_throws_made', 'free_throws_attempted',
                 'offensive_rebounds', 'defensive_rebounds', 'total_rebounds',
                 'assists', 'turnovers', 'steals', 'blocks', 'personal_fouls',
                 'fantasy_points', 'double_doubles', 'triple_doubles', 'possessions'}


class PlayerRow(BaseModel):
    model_config = ConfigDict(extra='forbid')
    player_name: str = Field(min_length=1)
    player_url: HttpUrl
    team_code: str = Field(min_length=2, max_length=3)
    team_name: str = Field(min_length=1)
    team_url: HttpUrl
    age: int = Field(ge=18)


StatRow = create_model(
    'StatRow',
    __config__=ConfigDict(extra='forbid', strict=True),
    **{name: (int if name in INTEGER_STATS else float,
              Field(ge=0) if name in INTEGER_STATS else Field())
       for name in STAT_COLUMNS.values()},
)


def hyperlink(cell) -> str:
    target = cell.hyperlink.target if cell.hyperlink else None
    if not target or urlparse(target).scheme not in ('http', 'https'):
        raise ValueError(f'Hyperlien HTTP(S) manquant dans {cell.coordinate}')
    return target


def load_rows(path: Path):
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = book['Données NBA']
        expected = ['Player', 'Team', 'Age', 'GP', 'W', 'L', 'Min', 'PTS']
        if [sheet.cell(2, i).value for i in range(1, 9)] != expected:
            raise ValueError('En-têtes inattendus : vérifier la ligne 2 de « Données NBA ».')
        teams = {str(row[0].value).strip(): str(row[1].value).strip()
                 for row in book['Equipe'].iter_rows(min_row=2, max_col=2)
                 if row[0].value and row[1].value}
        rows = []
        for row in sheet.iter_rows(min_row=3, max_col=45):
            if not row[0].value:
                continue
            try:
                code = str(row[1].value).strip()
                player = PlayerRow.model_validate({
                    'player_name': str(row[0].value).strip(),
                    'player_url': hyperlink(row[0]),
                    'team_code': code,
                    'team_name': teams[code],
                    'team_url': hyperlink(row[1]),
                    'age': row[2].value,
                })
                raw = {name: row[index - 1].value for index, name in STAT_COLUMNS.items()}
                for name, value in raw.items():
                    if value is None or isinstance(value, bool) or not isinstance(value, (int, float)):
                        raise ValueError(f'{name} doit être numérique : {value!r}')
                    if name in INTEGER_STATS and int(value) != value:
                        raise ValueError(f'{name} doit être un entier : {value!r}')
                stat = StatRow.model_validate({
                    name: int(value) if name in INTEGER_STATS else float(value)
                    for name, value in raw.items()
                })
                if stat.wins + stat.losses != stat.games_played:
                    raise ValueError('W + L différent de GP')
                rows.append((player, stat.model_dump()))
            except (KeyError, ValueError, ValidationError) as exc:
                raise ValueError(f'Ligne Excel {row[0].row} : {exc}') from exc
        if not rows:
            raise ValueError('Aucun joueur trouvé.')
        names = [p.player_name for p, _ in rows]
        if len(names) != len(set(names)):
            raise ValueError('Noms de joueurs en double : le schéma actuel impose UNIQUE(player_name).')
        return rows
    finally:
        book.close()


def import_excel(source: Path, schema: Path, db: Path, season: str, season_type: str):
    rows = load_rows(source)  # Valider l'intégralité du fichier avant toute écriture.
    if not schema.is_file():
        raise FileNotFoundError(schema)
    db.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db) as conn:
        conn.execute('PRAGMA foreign_keys = ON')
        conn.executescript(schema.read_text(encoding='utf-8'))
        actual = {r[1] for r in conn.execute('PRAGMA table_info(stats)')}
        missing = set(STAT_COLUMNS.values()) - actual
        if missing:
            raise ValueError(f'Colonnes absentes de stats : {sorted(missing)}')
        if conn.execute('SELECT 1 FROM reports WHERE source_file=? AND season=? AND season_type=?',
                        (source.name, season, season_type)).fetchone():
            raise ValueError('Ce fichier et cette saison ont déjà été importés.')
        try:
            conn.execute('BEGIN')
            cursor = conn.execute(
                'INSERT INTO reports (season, season_type, source_file, description) VALUES (?, ?, ?, ?)',
                (season, season_type, source.name, 'Statistiques agrégées par joueur ; aucun match individuel.'))
            report_id = cursor.lastrowid
            columns = list(STAT_COLUMNS.values())
            placeholders = ', '.join('?' for _ in columns)
            stat_sql = f"INSERT INTO stats (player_id, report_id, {', '.join(columns)}) VALUES (?, ?, {placeholders})"
            for player, values in rows:
                p = player.model_dump(mode='json')
                cursor = conn.execute(
                    'INSERT INTO players (player_name, player_url, team_code, team_name, team_url, age) '
                    'VALUES (:player_name, :player_url, :team_code, :team_name, :team_url, :age)', p)
                conn.execute(stat_sql, (cursor.lastrowid, report_id, *(values[c] for c in columns)))
            assert conn.execute('PRAGMA foreign_key_check').fetchall() == []
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    print(f'Import terminé : {len(rows)} joueurs et statistiques, 1 rapport. Base : {db}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('excel', type=Path)
    parser.add_argument('--schema', type=Path, default=Path('database/schema.sql'))
    parser.add_argument('--db', type=Path, default=Path('data/nba.sqlite'))
    parser.add_argument('--season', required=True, help='Saison vérifiée auprès de la source, ex. 2024-25')
    parser.add_argument('--season-type', default='Regular Season')
    args = parser.parse_args()
    import_excel(args.excel, args.schema, args.db, args.season, args.season_type)


if __name__ == '__main__':
    main()